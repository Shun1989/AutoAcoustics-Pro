"""Chinese DOCX reports from immutable results, never a second analysis path."""
from __future__ import annotations

from dataclasses import dataclass, fields
from datetime import datetime
from io import BytesIO
import hashlib
import json
import math
import os
from pathlib import Path
import re
import tempfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import numpy as np

from .batch import check_cancel, ordered_metrics, valid_metric_value, INVALID_METRIC_STATUSES
from .model import (AnalysisResult, BatchResult, ComparisonResult, MeasurementContext,
                    ValidationError, to_plain)

STATUS_NAMES = {'ok':'有效', 'complete':'有效', 'silent':'静音', 'estimated':'估计',
                'uncalibrated':'未校准', 'failed':'失败', 'not_applicable':'不适用',
                'partial':'部分指标失败',
                'cancelled':'已取消', 'running':'运行中', 'disabled':'未启用',
                'unavailable':'不可用', 'unknown':'未知'}
METRIC_NAMES = {'LZeq':'平均 Z 声级 LZeq', 'LAeq':'平均 A 声级 LAeq',
                'LAFmax':'最大 Fast A 声级 LAFmax', 'LASmax':'最大 Slow A 声级 LASmax',
                'Nmean':'时变响度时间平均 Nmean', 'Nmax':'时变响度最大值 Nmax',
                'Nstationary':'稳态响度 Nstationary', 'RMS':'均方根 RMS',
                'dBFS':'数字 RMS dBFS', 'RMS_dBFS':'数字 RMS dBFS', 'duration':'事件时长'}
CONTEXT_NAMES = dict(zip((field.name for field in fields(MeasurementContext)), (
    '样件', '样件版本', '测量层级', '机构', '供电', '供电来源', '负载', '负载布局',
    '工装', '紧固方式', '环境', '温度', '麦克风测点', '麦克风距离', '麦克风方向',
    '背景', '实际运动方向', '电机旋向', '旋向观察基准', '行程', '循环', '重复次数', '备注')))


@dataclass(frozen=True)
class ReportSpec:
    result: AnalysisResult | BatchResult | ComparisonResult
    project_name: str = ''
    project_version: str | int | None = None
    operator: str = ''
    notes: str = ''
    cancel: object = None


def _text(value, unknown='未知'):
    if value is None or value == '' or value == 'unknown':
        return unknown
    return str(value)


def _font(run, size=11, bold=False, color='000000'):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')


def _style(document):
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(1)
    for name, size in [('Normal',11), ('Title',21), ('Heading 1',15), ('Heading 2',12), ('Caption',10)]:
        style = document.styles[name]
        style.font.name = 'Arial'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        if name != 'Normal':
            style.paragraph_format.keep_with_next = True
        if name == 'Title':
            for border in style._element.findall('.//' + qn('w:pBdr')):
                border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(footer.add_run('AutoAcoustics Pro  ·  '), 9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
    footer._p.append(field)
    document.core_properties.title = 'AutoAcoustics Pro 声学测量结果'
    document.core_properties.author = 'AutoAcoustics Pro'
    document.core_properties.subject = '声学指标与测量配置可追溯报告'


def _paragraph(document, value, style=None):
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(str(value))
    _font(run, 10 if style == 'Caption' else 11)
    return paragraph


def _table(document, headers, rows, widths, left_columns=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    border = OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        element = OxmlElement(f'w:{name}')
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:
            element.set(qn(f'w:{key}'),value)
        border.append(element)
    table._tbl.tblPr.append(border)
    header = table.rows[0]
    repeat = OxmlElement('w:tblHeader'); header._tr.get_or_add_trPr().append(repeat)
    records = [headers] + list(rows)
    for index, record in enumerate(records):
        row = header if index == 0 else table.add_row()
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for column,(cell,value,width) in enumerate(zip(row.cells,record,widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                side_element = OxmlElement(f'w:{side}')
                side_element.set(qn('w:w'),'90'); side_element.set(qn('w:type'),'dxa')
                margins.append(side_element)
            properties.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'),'23415D' if index == 0 else 'F1F5F8' if index % 2 == 0 else 'FFFFFF')
            properties.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.1
            if index == 0:
                paragraph.paragraph_format.keep_with_next = True
            left_columns = (0,len(headers)-1) if left_columns is None else left_columns
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column in left_columns else WD_ALIGN_PARAGRAPH.CENTER
            _font(paragraph.add_run(str(value)), 10, index == 0, 'FFFFFF' if index == 0 else '000000')
    document.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def _paired_table(document, pairs):
    pairs=list(pairs)
    records=[]
    for start in range(0,len(pairs),2):
        left=pairs[start]
        right=pairs[start+1] if start+1<len(pairs) else ('','')
        records.append((*left,*right))
    return _table(document,['字段','记录','字段','记录'],records,
                  [1.1,2.15,1.1,2.15],left_columns=(0,1,2,3))


def _metric_table(document, metrics, method_reference=None):
    rows = []
    for name in ordered_metrics(metrics):
        metric = metrics[name]
        value = valid_metric_value(metric)
        display = '—' if value is None else '−∞' if value == -math.inf else format(value,'.10g')
        method=f'{method_reference} / {name}' if method_reference and metric.method else metric.method
        explanation = '\n'.join(part for part in (method, metric.message) if part) or '方法未记录'
        status = STATUS_NAMES.get(metric.status,metric.status)
        if metric.value is not None and value is None and metric.status not in INVALID_METRIC_STATUSES:
            status = '失败'; explanation += '\n指标为无效非有限值'
        rows.append((METRIC_NAMES.get(name,name),display,metric.unit,status,explanation))
    if rows:
        widths=[1.9,1.25,1.1,.65,1.6] if method_reference else [1.2,1.3,.85,.65,2.5]
        _table(document,['指标','数值','单位','状态','方法及说明'],rows,widths)
    else:
        _paragraph(document,'本条没有已计算的指标。')


def _context(document, context):
    document.add_heading('测量工况', level=2)
    _paired_table(document,[(CONTEXT_NAMES[name],_text(value))
                           for name,value in to_plain(context).items() if name!='notes'])
    _paragraph(document,'备注 '+_text(context.notes,'未填写'))
    _paragraph(document,'未知工况会限制不同结果之间的可比性，报告不自动判断优劣或故障。')


def _configuration(document, result):
    settings, provenance = result.settings, result.provenance
    end = settings.end
    if end is None:
        end = provenance.get('event_end_sample',provenance.get('end_sample',
                             provenance.get('source_frames',provenance.get('frames'))))
    rate = provenance.get('sample_rate')
    source = result.calibration.source if result.calibration else provenance.get('calibration_source',
                                        provenance.get('calibration',{}).get('source'))
    rows = [('结果 ID',result.result_id), ('输入 SHA256',result.input_hash or '未记录'),
            ('源文件',str(result.path) if result.path else '未记录'), ('通道索引',settings.channel),
            ('原样本范围',f'[{settings.start}, {end})' if end is not None else f'[{settings.start}, 文件末尾)'),
            ('采样率',f'{rate} Hz' if rate else '未记录'), ('事件',settings.event_label),
            ('输入单位',_text(provenance.get('source_unit',provenance.get('input_unit')),'未记录')),
            ('分析单位',_text(provenance.get('unit',provenance.get('input_unit')),'未记录')),
            ('校准来源',_text(source,'未校准或来源未记录')),
            ('响度声场','自由场' if settings.field_type == 'free' else '扩散场'),
            ('去直流','是' if settings.remove_dc else '否'),
            ('FFT 请求设置',f'{settings.fft_size} 点  {settings.window} 窗'),
            ('STFT 参数',f'{settings.stft_size} 点  步长 {settings.stft_hop} 点'),
            ('SPL 时间计权',settings.spl_time_weighting),
            ('时变响度','启用 ISO 532-1 Zwicker' if settings.compute_loudness else '未启用'),
            ('额外稳态响度','启用' if settings.stationary_loudness else '未启用')]
    if rate and end is not None:
        begin_time,end_time=_event_times(settings,provenance,end)
        rows.append(('原录音时间范围',f'[{begin_time:.9g}, {end_time:.9g}) s'))
    _paired_table(document,rows)
    if result.calibration:
        profile = result.calibration
        _paragraph(document,f'校准配置 {profile.name}；输入单位 {profile.input_unit}；换算系数 {profile.coefficient:.12g}；记录日期 {_text(profile.date,"未记录")}。')


def _provenance(document, result):
    document.add_heading('算法和来源记录',level=2)
    _provenance_values(document,result.provenance)
    if not result.provenance:
        _paragraph(document,'算法版本和来源尚未记录。')


def _provenance_values(document, values):
    for key,value in values.items():
        plain = to_plain(value)
        if key == 'source_metadata':
            _paragraph(document,'source_metadata SHA256  '+_metadata_digest(plain))
            _paired_table(document,[(key,_text(value)) for key,value in _source_metadata_summary(plain).items()])
            _paragraph(document,'完整原始头与传感器 XML 保留在结果 provenance 中，可随 CSV provenance_json 和项目来源快照保存；这些字段用于来源复核，未据此再次校准。')
            continue
        if key == 'calibration' and isinstance(plain,dict):
            hashes=plain.get('applicable_hashes')
            if isinstance(hashes,list) and len(hashes)>8:
                # The current source SHA is printed in the configuration. The
                # whole immutable manifest remains in the result and CSV;
                # its exact canonical digest identifies it without pages of
                # repeated hashes in every batch record.
                encoded=json.dumps(hashes,ensure_ascii=False,sort_keys=True,
                                   separators=(',',':')).encode('utf-8')
                plain=dict(plain)
                plain['applicable_hashes']={'记录':f'{len(hashes)} 项',
                    '完整清单 SHA256':hashlib.sha256(encoded).hexdigest(),
                    '说明':'完整清单保存在原结果及 CSV 校准记录，本条输入 SHA256 见测量配置'}
        text = json.dumps(plain, ensure_ascii=False, sort_keys=True) if isinstance(plain,(dict,list)) else str(plain)
        paragraph=_paragraph(document,f'{key}  {text}')
        _font(paragraph.runs[0],10)
        paragraph.paragraph_format.space_after=Pt(3)


def _metadata_digest(metadata):
    encoded=json.dumps(to_plain(metadata),ensure_ascii=False,sort_keys=True,
                       separators=(',',':')).encode('utf-8')
    return hashlib.sha256(encoded).hexdigest()


def _source_metadata_summary(metadata):
    """Known source and sensor attributes only; raw structures stay in CSV/project."""
    metadata=to_plain(metadata)
    if not isinstance(metadata,dict):
        return {'来源摘要':'没有可解析的结构化字段'}
    summary={key:metadata[key] for key in ('format','subtype','bit_depth','source_dtype',
             'normalization','lossy','source_quantity','calibration_applied_by_importer')
             if key in metadata and not isinstance(metadata[key],(dict,list))}
    head=metadata.get('head',{})
    if not isinstance(head,dict):head={}
    summary.update({'HEAD '+key:head[key] for key in ('version','release','byte_order',
                    'implementation_type','text_encoding') if key in head})
    sensor=metadata.get('sensor',{})
    if not isinstance(sensor,dict):sensor={}
    if not sensor:
        raw=head.get('comment_metadata','')
        if isinstance(raw,str) and not re.search(r'<!\s*(DOCTYPE|ENTITY)\b',raw,re.I):
            opening=re.search(r'<SensorInfo\b[^>]*>',raw)
            if opening:
                tag=opening.group()
                if not tag.endswith('/>'):tag=tag[:-1]+'/>'
                try:sensor=dict(ET.fromstring(tag).attrib)
                except ET.ParseError:summary['传感器摘要']='属性不能安全解析，见完整来源快照'
    for key in ('Name','Manufacturer','SerialNumber','SensorType','Model','OutputUnit',
                'Sensitivity','CalibrationFactor','CalibrationDate','Index','ChName'):
        if key in sensor:
            # Sensitivity/CalibrationFactor remain raw attributes. An output
            # unit alone is not a validated physical sensitivity unit.
            label='Sensor '+key+(' 原属性未用于重校准' if key in ('Sensitivity','CalibrationFactor') else '')
            summary[label]=sensor[key]
    return summary or {'来源摘要':'关键来源字段未记录'}


def _event_times(settings,provenance,end):
    rate=provenance.get('sample_rate')
    origin=provenance.get('time_origin_seconds',0.)
    return (provenance.get('event_start_time_seconds',origin+settings.start/rate),
            provenance.get('event_end_time_seconds',origin+end/rate))


def _figure(title, ylabel, xlabel):
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib import font_manager
    figure = Figure(figsize=(6.4,2.8),dpi=150,layout='constrained')
    FigureCanvasAgg(figure)
    ax = figure.add_subplot(111)
    font = font_manager.FontProperties(family='Microsoft YaHei')
    ax.set_title(title,fontproperties=font,fontsize=11)
    ax.set_xlabel(xlabel,fontproperties=font,fontsize=10)
    ax.set_ylabel(ylabel,fontproperties=font,fontsize=10)
    ax.grid(True,alpha=.22)
    for label in ax.get_xticklabels()+ax.get_yticklabels():
        label.set_fontproperties(font)
    return figure, ax, font


def _insert_figure(document, figure, caption):
    image = BytesIO()
    figure.savefig(image,format='png',dpi=150)
    image.seek(0)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(image,width=Inches(6.4))
    _paragraph(document,caption,style='Caption').paragraph_format.keep_with_next = False
    figure.clear()


def _decimate(x,y,limit=10000):
    """Display-only extrema preservation, never metric recomputation."""
    if len(y) <= limit:
        return x,y
    edges = np.linspace(0,len(y),limit // 2 + 1,dtype=int)
    indices=[]
    for left,right in zip(edges[:-1],edges[1:]):
        block=y[left:right]
        indices.extend(sorted((left+int(np.argmin(block)),left+int(np.argmax(block)))))
    indices=np.asarray(indices)
    return x[indices],y[indices]


def _curves(result):
    if result.detail is None:
        return []
    arrays,units=result.detail.arrays,result.detail.units
    curves=[]
    wave_key='wave_time' if 'wave_time' in arrays else 'waveform_time'
    if wave_key not in arrays and 'waveform' in arrays and result.provenance.get('sample_rate'):
        wave_x=result.provenance.get('time_origin_seconds',0.)+(np.arange(len(arrays['waveform']))+result.settings.start)/result.provenance['sample_rate']
    else:
        wave_x=arrays.get(wave_key)
    candidates=[('波形','时间 / s',wave_x,'waveform'),
        ('幅度谱','频率 / Hz',arrays.get('frequency'),'spectrum_db'),
        ('功率谱密度','频率 / Hz',arrays.get('psd_frequency',arrays.get('frequency')),'psd'),
        ('三分之一倍频程','精确中心频率 / Hz',arrays.get('octave_centers'),'octave_levels'),
        ('Fast A 声级历程','原录音时间 / s',arrays.get('spl_time'),'laf'),
        ('Slow A 声级历程','原录音时间 / s',arrays.get('spl_time'),'las'),
        ('时变响度历程','原录音时间 / s',arrays.get('loudness_time'),'loudness')]
    for title,xlabel,x,ykey in candidates:
        if x is not None and ykey in arrays:
            x,y=np.asarray(x),np.asarray(arrays[ykey])
            if x.ndim==y.ndim==1 and len(x)==len(y) and len(x):
                curves.append((title,xlabel,x,y,units.get(ykey,'单位未记录'),ykey))
    return curves


def _line(ax,x,y,label=None):
    mask=np.isfinite(x)&np.isfinite(y)
    if mask.any():
        x,y=_decimate(x[mask],y[mask])
        ax.plot(x,y,linewidth=.9,label=label)
    else:
        ax.text(.5,.5,'无有限值曲线',transform=ax.transAxes,ha='center',
                fontfamily='Microsoft YaHei')


def _single_charts(document,result,cancel):
    if result.detail is None:
        _paragraph(document,'仅有摘要，未保存分析详情。报告没有重算波形或编造历程曲线。')
        return
    document.add_heading('已有分析图表',level=2)
    for title,xlabel,x,y,unit,_ in _curves(result):
        check_cancel(cancel)
        figure,ax,_=_figure(title,unit,xlabel)
        _line(ax,x,y)
        _insert_figure(document,figure,f'{title}，单位 {unit}，读取已保存的分析详情。')
    arrays=result.detail.arrays
    if all(key in arrays for key in ('stft_time','stft_frequency','stft_db')):
        data=np.asarray(arrays['stft_db']); times=np.asarray(arrays['stft_time']); frequencies=np.asarray(arrays['stft_frequency'])
        if data.shape==(len(frequencies),len(times)) and data.size:
            if not np.isfinite(data).any():
                _paragraph(document,'语谱图没有有限的 dB 数值，例如静音负无穷；未绘制虚构色阶。')
                return
            check_cancel(cancel)
            figure,ax,font=_figure('语谱图','频率 / Hz','原录音时间 / s')
            # Bound raster display resolution. Averaging or recalculating
            # scientific spectra is not performed for the report.
            tstep=max(1,math.ceil(len(times)/1200)); fstep=max(1,math.ceil(len(frequencies)/600))
            raster=np.ma.masked_invalid(data[::fstep,::tstep])
            image=ax.pcolormesh(times[::tstep],frequencies[::fstep],raster,shading='nearest',cmap='viridis')
            colorbar=figure.colorbar(image,ax=ax)
            colorbar.set_label(result.detail.units.get('stft_db','dB'),fontproperties=font,fontsize=9)
            _insert_figure(document,figure,'语谱图读取已有 STFT 数值，绘图尺寸限制只用于展示。')


def _single(document,result,spec,charts=True):
    document.add_heading('测量配置',level=1)
    _configuration(document,result)
    document.add_heading('指标结果',level=1)
    _metric_table(document,result.summary.metrics)
    for warning in result.summary.warnings:
        _paragraph(document,f'质量或适用性说明  {warning}')
    _context(document,result.context)
    if charts:
        _single_charts(document,result,spec.cancel)
    _provenance(document,result)


def _batch_charts(document,batch,spec):
    for metric_name in ('LZeq','LAeq','Nmean','Nmax'):
        values=[]; labels=[]; unit=None
        for index,row in enumerate(batch.items,1):
            if row.status not in ('ok','complete','partial') or row.result is None:
                continue
            metric=row.result.summary.metrics.get(metric_name)
            value=valid_metric_value(metric) if metric else None
            if value is not None and math.isfinite(value):
                if unit is not None and unit!=metric.unit:
                    continue
                unit=metric.unit; values.append(value); labels.append(str(index))
        if values:
            check_cancel(spec.cancel)
            title=METRIC_NAMES[metric_name]+' 批量摘要'
            figure,ax,_=_figure(title,unit,'分析条目序号')
            ax.plot(np.asarray(labels,dtype=int),values,'o-',markersize=3,linewidth=.7)
            _insert_figure(document,figure,'仅汇总具有有限有效值的条目；失败、未校准、不适用和静音负无穷没有被填成零。')


def _shared_method(value):
    """Separate event coordinates, which are printed on every item, from methods."""
    plain=to_plain(value)
    if isinstance(plain,dict):
        return {key:_shared_method(item) for key,item in plain.items()
                if key not in ('event_start_sample','event_end_sample')}
    if isinstance(plain,list):
        return [_shared_method(item) for item in plain]
    return plain


def _batch_records(document,batch):
    groups={name:[] for name in ('S','P','M','H','Q')}
    indices={name:{} for name in groups}
    references=[]
    contexts=[to_plain(row.result.context if row.result else row.specification.context)
              for row in batch.items if row.result or row.specification]
    common={key:value for key,value in contexts[0].items()
            if all(context.get(key)==value for context in contexts)} if contexts else {}
    def register(prefix,value):
        plain=to_plain(value)
        identity=json.dumps(plain,ensure_ascii=False,sort_keys=True,separators=(',',':'))
        if identity not in indices[prefix]:
            indices[prefix][identity]=f'{prefix}{len(groups[prefix])+1}'
            groups[prefix].append((indices[prefix][identity],value))
        return indices[prefix][identity]
    for row in batch.items:
        result=row.result;item=row.specification
        settings=result.settings if result else item.settings if item else None
        context=result.context if result else item.context if item else None
        profile=result.calibration if result else item.profile if item else None
        provenance=result.provenance if result else {}
        if settings is None:
            references.append({});continue
        processing={key:value for key,value in to_plain(settings).items()
                    if key not in ('channel','start','end','event_label')}
        calibration=to_plain(profile) if profile else to_plain(provenance.get('calibration'))
        methods={key:_shared_method(provenance[key]) for key in ('versions','method_revision','methods')
                 if key in provenance}
        if result:
            methods['metric_methods']={key:metric.method for key,metric in result.summary.metrics.items()}
        context_plain=to_plain(context) if context is not None else {}
        references.append({'S':register('S',processing),'C':'C1' if context is not None else '未记录',
            'context_differences':{key:value for key,value in context_plain.items() if key not in common},
            'P':register('P',calibration),'M':register('M',methods) if result else '未生成',
            'H':register('H',_source_metadata_summary(provenance['source_metadata'])) if 'source_metadata' in provenance else '未记录',
            'Q':register('Q',{'source_flags':to_plain(provenance.get('quality',[])),
                             'warnings':to_plain(result.summary.warnings)}) if result else '未生成'})
    document.add_heading('共同处理和测量记录',level=1)
    _paragraph(document,'相同的处理设置、校准、算法和来源摘要只列一次。工况 C1 只列所有已提供工况条目完全相同的字段；其余字段在各条列出，不表示不同样件是同一次试验。没有工况记录的条目不套用 C1。每条指标仍来自各自结果，完整来源快照保存在 CSV 和项目中。')
    names={'fft_size':'FFT 请求点数','window':'窗函数','stft_size':'STFT 点数',
           'stft_hop':'STFT 步长 点','spl_time_weighting':'SPL 时间计权',
           'history_step':'SPL 历程步长 秒','field_type':'响度声场',
           'remove_dc':'去直流','compute_loudness':'时变响度',
           'stationary_loudness':'额外稳态响度'}
    for identity,value in groups['S']:
        document.add_heading('处理记录 '+identity,level=2)
        _paired_table(document,[(names.get(key,key),
            '自由场' if key=='field_type' and item=='free' else
            '扩散场' if key=='field_type' and item=='diffuse' else
            '是' if item is True else '否' if item is False else _text(item))
            for key,item in value.items()])
    if contexts:
        document.add_heading('工况记录 C1 共同字段',level=2)
        if common:
            _paired_table(document,[(CONTEXT_NAMES.get(key,key),_text(value,'未填写' if key=='notes' else '未知'))
                                   for key,value in common.items()])
        else:_paragraph(document,'没有完全相同的工况字段，全部记录在各条差异中。')
        _paragraph(document,'未知工况保持未知，会限制可比性；没有自动判断优劣或故障。')
    for identity,value in groups['P']:
        document.add_heading('校准记录 '+identity,level=2)
        _provenance_values(document,{'calibration':value})
    for identity,value in groups['M']:
        document.add_heading('方法记录 '+identity,level=2)
        _provenance_values(document,{key:item for key,item in value.items() if key!='metric_methods'})
        _table(document,['指标','指标方法'],[(METRIC_NAMES.get(key,key),item or '方法未记录')
                 for key,item in value.get('metric_methods',{}).items()],[2.2,4.3],left_columns=(0,1))
    for identity,value in groups['H']:
        document.add_heading('来源摘要 '+identity,level=2)
        _paired_table(document,[(key,_text(item)) for key,item in value.items()])
    if groups['H']:
        _paragraph(document,'每条打印完整 source_metadata 的规范 JSON SHA256，引用这些关键来源及传感器字段摘要。原始头与传感器 XML 只保存在 CSV provenance_json 和项目快照中；未依据原始属性再次校准。')
    for identity,value in groups['Q']:
        document.add_heading('质量记录 '+identity,level=2)
        _paragraph(document,'来源质量标记 '+('；'.join(map(str,value['source_flags'])) or '无已记录标记，不代表全部前端状态已知'))
        for warning in value['warnings']:
            _paragraph(document,'质量或适用性说明 '+str(warning))
        if not value['warnings']:_paragraph(document,'分析结果未记录额外警告。')
    return references


def _batch_item_configuration(document,row,references):
    result=row.result;item=row.specification
    settings=result.settings if result else item.settings if item else None
    provenance=result.provenance if result else {}
    if settings is None:
        _paragraph(document,'原条目配置缺失，无法复核通道、事件或工况。');return
    end=settings.end
    if end is None:end=provenance.get('event_end_sample',provenance.get('source_frames'))
    source=provenance.get('source_unit',provenance.get('input_unit'))
    input_hash=result.input_hash if result else item.input_hash
    _paragraph(document,'结果 ID '+(result.result_id if result else '未生成'))
    _paragraph(document,'输入 SHA256 '+(input_hash or '未验证'))
    _paragraph(document,'源文件 '+str(row.path))
    metadata=to_plain(provenance.get('source_metadata',{}))
    if not isinstance(metadata,dict):metadata={}
    pairs=[('通道索引',settings.channel),('通道名',_text(provenance.get('channel_name'),'未验证')),
           ('事件',settings.event_label),('源格式',_text(metadata.get('format'),'未验证')),
           ('原样本范围',f'[{settings.start}, {end})' if end is not None else f'[{settings.start}, 文件末尾 未验证)'),
           ('原采样率',f'{provenance["sample_rate"]} Hz' if provenance.get('sample_rate') else '未验证'),
           ('输入／分析单位',f'{_text(source,"未验证")} / {_text(provenance.get("unit"),"未验证")}'),
           ('处理记录',references.get('S','未记录')),('工况记录',references.get('C','未记录')),
           ('校准记录',references.get('P','未记录')),('方法记录',references.get('M','未生成')),
           ('来源摘要',references.get('H','未记录')),('质量记录',references.get('Q','未生成'))]
    preprocessing=provenance.get('preprocessing',{})
    if preprocessing:
        pairs.append(('直流处理实录',f'去直流 {preprocessing.get("remove_dc","未记录")}；减去 {preprocessing.get("removed_dc","未记录")} {_text(provenance.get("unit"),"单位未记录")}；范围 {preprocessing.get("scope","未记录")}'))
    if end is not None and provenance.get('sample_rate'):
        begin_time,end_time=_event_times(settings,provenance,end)
        pairs.extend([('源时间范围',f'[{begin_time:.9g}, {end_time:.9g}) s'),
                      ('源时钟起点',f'{provenance.get("time_origin_seconds",0.):.9g} s')])
    _paired_table(document,pairs)
    differences=references.get('context_differences',{})
    if differences:
        _paragraph(document,'本条工况差异字段（其余字段见 C1）')
        _paired_table(document,[(CONTEXT_NAMES.get(key,key),_text(value,'未填写' if key=='notes' else '未知'))
                              for key,value in differences.items()])
    if result:
        if 'source_metadata' in provenance:
            _paragraph(document,'source_metadata SHA256 '+_metadata_digest(provenance['source_metadata']))


def _batch(document,batch,spec):
    counts={status:sum(row.status==status for row in batch.items) for status in ('ok','complete','partial','failed','cancelled')}
    _paragraph(document,f'{len({str(row.path) for row in batch.items})} 个源文件，{len(batch.items)} 个独立分析条目；完成 {counts["ok"]+counts["complete"]} 条，部分指标失败 {counts["partial"]} 条，失败 {counts["failed"]} 条，取消 {counts["cancelled"]} 条。')
    _paragraph(document,'批处理仅有摘要，未常驻或保存逐条完整绘图数组；以下完整指标来自已有结果。')
    document.add_heading('运行清单',level=1)
    _table(document,['序号','文件及条目 ID','状态','错误'],
           [(index,f'{row.path}\n{row.item_id}',STATUS_NAMES.get(row.status,row.status),row.error or '—')
            for index,row in enumerate(batch.items,1)], [.45,3.0,.65,2.4])
    document.add_heading('批量指标摘要图',level=1)
    _batch_charts(document,batch,spec)
    references=_batch_records(document,batch)
    for index,row in enumerate(batch.items,1):
        check_cancel(spec.cancel)
        document.add_heading(f'条目 {index} 测量记录',level=1)
        _paragraph(document,f'分析条目 ID {row.item_id}；状态 {STATUS_NAMES.get(row.status,row.status)}。')
        if row.error:
            _paragraph(document,f'错误说明 {row.error}')
        _batch_item_configuration(document,row,references[index-1])
        if row.result:
            _metric_table(document,row.result.summary.metrics,references[index-1].get('M'))
        elif row.specification:
            _paragraph(document,'未完成分析，未生成指标。原处理、校准和工况配置见所引记录。')
        else:
            _paragraph(document,'未完成分析，未生成指标。')


def _comparison(document,result,spec):
    if result.direction!='right_minus_left':
        raise ValidationError('比较报告仅支持明确的右侧减左侧差值。')
    _paragraph(document,'差值方向为右侧减左侧。dB 差值和 sone 差值分别保留原单位；缺失、静音和失败结果不产生虚假的百分比或优劣结论。')
    document.add_heading('条件差异和可比性',level=1)
    for condition in result.condition_differences:
        _paragraph(document,condition)
    if not result.condition_differences:
        _paragraph(document,'已记录字段未发现差异；未知工况仍限制可比性。')
    document.add_heading('右侧减左侧指标差值',level=1)
    _metric_table(document,result.differences)
    document.add_heading('双结果已有曲线',level=1)
    left_curves={item[-1]:item for item in _curves(result.left)}
    right_curves={item[-1]:item for item in _curves(result.right)}
    if not left_curves or not right_curves:
        _paragraph(document,'至少一侧仅有摘要，缺少对应详情，未绘制不存在的双结果曲线。')
    for key in left_curves:
        if key not in right_curves:
            continue
        check_cancel(spec.cancel)
        left,right=left_curves[key],right_curves[key]
        if left[4]!=right[4]:
            _paragraph(document,f'{left[0]}单位不同 左 {left[4]} 右 {right[4]}，不共用纵轴。')
            continue
        figure,ax,font=_figure(left[0]+' 双结果叠图',left[4],left[1])
        _line(ax,left[2],left[3],'左侧'); _line(ax,right[2],right[3],'右侧')
        ax.legend(prop=font,fontsize=9)
        _insert_figure(document,figure,'两侧使用各自原有坐标轴；不同频率网格仅叠图，未按数组位置相减。')
    for label,item in [('左侧',result.left),('右侧',result.right)]:
        heading=document.add_heading(label+'结果和配置',level=1)
        heading.paragraph_format.page_break_before=True
        _single(document,item,spec,charts=False)


def _save_document(document,path):
    document.save(path)


def write_docx(result: AnalysisResult | BatchResult | ComparisonResult | ReportSpec,
               destination: Path) -> Path:
    spec=result if isinstance(result,ReportSpec) else ReportSpec(result)
    if not isinstance(spec.result,(AnalysisResult,BatchResult,ComparisonResult)):
        raise ValidationError('DOCX 报告必须来自单文件、批处理或比较结果。')
    check_cancel(spec.cancel)
    document=Document(); _style(document)
    title='批量声学分析报告' if isinstance(spec.result,BatchResult) else '声学结果比较报告' if isinstance(spec.result,ComparisonResult) else '声学分析报告'
    document.add_paragraph(title,style='Title')
    _paragraph(document,'本报告记录既有声学分析结果、测量配置、工况和适用性，供座椅电机试验复核。指标来自同一份分析结果，没有在导出时重新计算。')
    _paragraph(document,f'生成时间 {datetime.now().astimezone().isoformat(timespec="seconds")}；项目 {_text(spec.project_name,"未指定")}；项目版本 {_text(spec.project_version,"未记录")}；操作者 {_text(spec.operator,"未记录")}。')
    if spec.notes:
        _paragraph(document,spec.notes)
    if isinstance(spec.result,AnalysisResult):
        _single(document,spec.result,spec)
    elif isinstance(spec.result,BatchResult):
        _batch(document,spec.result,spec)
    else:
        _comparison(document,spec.result,spec)
    check_cancel(spec.cancel)
    destination=Path(destination)
    destination.parent.mkdir(parents=True,exist_ok=True)
    descriptor,name=tempfile.mkstemp(prefix=f'.{destination.stem}.',suffix='.tmp.docx',dir=destination.parent)
    os.close(descriptor); temporary=Path(name)
    try:
        _save_document(document,temporary)
        check_cancel(spec.cancel)
        temporary.replace(destination)
        return destination
    finally:
        if temporary.exists():
            temporary.unlink()
