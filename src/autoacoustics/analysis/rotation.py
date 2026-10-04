"""Traceable rotational evidence from audio, with optional measured RPM tracking.

Manual RPM is a constant-speed frequency conversion, never a tachometer estimate.
Measured RPM is piecewise-linear in absolute audio time. Its exact integral is
resampled on an even-angle grid after conservative time-domain anti-aliasing.
The outputs are observed candidates and verification steps, not a fault verdict.

Primary definitions: NI, Order Analysis Based on Resampling;
https://www.ni.com/docs/en-US/bundle/diadem/page/genmaths/genmaths/calc_oa_resampling.htm
and SciPy signal.resample_poly (zero-phase low-pass FIR resampling):
https://docs.scipy.org/doc/scipy/reference/generated/scipy.signal.resample_poly.html
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable
import base64
import csv
import html
import io
import json
import math

import numpy as np
from scipy.integrate import cumulative_trapezoid
from scipy.signal import butter, find_peaks, resample_poly, sosfiltfilt, stft

from .spectrum import compute_spectrum


@dataclass(frozen=True)
class RotationConfig:
    rpm: float = 3000.0
    gear_teeth: int = 0
    gear_ratio: float = 1.0
    pole_pairs: int = 0
    max_order: float = 50.0


@dataclass(frozen=True)
class RotationMarker:
    label: str
    frequency: float
    order: float
    frequency_low: float
    frequency_high: float


@dataclass(frozen=True)
class RotationPeak:
    frequency: float
    order: float
    level_db: float
    prominence_db: float
    frequency_kind: str = 'observed'


@dataclass(frozen=True)
class RotationFinding:
    mechanism: str
    evidence: tuple[str, ...]
    checks: tuple[str, ...]
    recommendations: tuple[str, ...]


@dataclass(frozen=True)
class RotationResult:
    frequency: np.ndarray
    spectrum_db: np.ndarray
    orders: np.ndarray
    order_db: np.ndarray
    map_times: np.ndarray
    map_orders: np.ndarray
    map_db: np.ndarray
    rpm_times: np.ndarray
    rpm_values: np.ndarray
    markers: tuple[RotationMarker, ...]
    peaks: tuple[RotationPeak, ...]
    findings: tuple[RotationFinding, ...]
    notes: tuple[str, ...]
    config: RotationConfig
    source_unit: str
    sample_rate: float
    time_origin: float
    duration: float
    rpm_mode: str
    amplitude_label: str
    effective_max_order: float
    samples_per_revolution: int
    frequency_reference_rpm: float


def _readonly(values):
    result = np.array(values, dtype=float, copy=True)
    result.setflags(write=False)
    return result


def _validate_config(config):
    if not isinstance(config, RotationConfig):
        raise ValueError('转频配置需要 RotationConfig。')
    for name in ('rpm', 'gear_ratio', 'max_order'):
        value = getattr(config, name)
        if not isinstance(value, (int, float, np.number)) or not np.isfinite(value) or value <= 0:
            raise ValueError(f'{name} 必须为有限正数。')
    for name in ('gear_teeth', 'pole_pairs'):
        value = getattr(config, name)
        if not isinstance(value, (int, float, np.number)) or not np.isfinite(value) or value < 0 or int(value) != value:
            raise ValueError(f'{name} 必须为非负整数；0 表示尚未提供。')


def _rpm_arrays(times, values):
    times = np.asarray(times, dtype=float)
    values = np.asarray(values, dtype=float)
    if times.ndim != 1 or values.ndim != 1 or len(times) != len(values) or len(times) < 2:
        raise ValueError('转速曲线需要至少两个一一对应的 time_s 与 rpm 样本。')
    if not np.isfinite(times).all() or not np.isfinite(values).all():
        raise ValueError('转速曲线时间与 rpm 必须为有限数值。')
    if np.any(np.diff(times) <= 0):
        raise ValueError('转速曲线 time_s 必须严格递增；不会自动排序或删除重复时间。')
    if np.any(values <= 0):
        raise ValueError('转速曲线 rpm 必须大于 0；请选择不包含停转/反转的区间。')
    return times, values


def load_rpm_csv(path: str | Path):
    """Read explicit seconds/RPM columns, preserving acquisition order.

    Accepted headings are time_s,rpm or 时间_s,转速_rpm. No unit guessing,
    extrapolation, sorting, blank-row interpolation or implicit millisecond conversion.
    """
    path = Path(path)
    text = path.read_text(encoding='utf-8-sig')
    reader = csv.DictReader(io.StringIO(text))
    headers = tuple((name or '').strip() for name in reader.fieldnames or ())
    time_key = next((name for name in ('time_s', '时间_s') if name in headers), None)
    rpm_key = next((name for name in ('rpm', '转速_rpm') if name in headers), None)
    if time_key is None or rpm_key is None:
        raise ValueError('CSV 必须包含 time_s,rpm（秒、每分钟转数）列；也可使用 时间_s,转速_rpm。')
    reader.fieldnames = list(headers)
    times, values = [], []
    for number, row in enumerate(reader, start=2):
        try:
            times.append(float(row[time_key]))
            values.append(float(row[rpm_key]))
        except (ValueError, TypeError, KeyError) as exc:
            raise ValueError(f'CSV 第 {number} 行的 time_s/rpm 不是有效数值。') from exc
    times, values = _rpm_arrays(times, values)
    return _readonly(times), _readonly(values)


def _phase_at(sample_times, rpm_times, rpm_values):
    # Piecewise linear RPM has a quadratic phase; integrate each CSV segment
    # exactly, including knots between audio samples (no sample-grid drift).
    cumulative = cumulative_trapezoid(rpm_values / 60.0, rpm_times, initial=0.0)
    indices = np.clip(np.searchsorted(rpm_times, sample_times, side='right')-1, 0, len(rpm_times)-2)
    delta = sample_times-rpm_times[indices]
    slope = np.diff(rpm_values)/np.diff(rpm_times)
    phase = cumulative[indices] + rpm_values[indices]*delta/60 + slope[indices]*delta**2/120
    return phase-phase[0]


def _db(amplitude):
    with np.errstate(divide='ignore'):
        return 20*np.log10(amplitude)


def _map(samples, sample_rate, window_samples, order_scale, time_origin=0., phase=None, original_times=None):
    count = len(samples)
    window_samples = min(count, max(32, int(window_samples)))
    hop = max(1, window_samples//4)
    # Bound heat-map width without changing its window or order resolution.
    hop = max(hop, math.ceil(max(0, count-window_samples)/511))
    hop = min(hop, window_samples)
    frequency, centres, values = stft(samples, fs=sample_rate, nperseg=window_samples,
                                    noverlap=window_samples-hop, detrend=False,
                                    boundary=None, padded=False, scaling='spectrum')
    factors = np.full(len(frequency), 2.)
    factors[0] = 1.
    if window_samples % 2 == 0:
        factors[-1] = 1.
    db = _db(abs(values)*factors[:, None])
    if phase is not None:
        centres = np.interp(centres, phase, original_times)
    else:
        centres = centres + time_origin
    return frequency/order_scale, centres, db


def _angular_signal(samples, sample_rate, absolute_times, phase, min_rpm, max_rpm, max_order):
    effective_max = min(float(max_order), .4*sample_rate/(max_rpm/60))
    if effective_max < .25:
        raise ValueError('采样率相对最高转速过低，无法可信地分析所需阶次。')
    ratio = max_rpm/min_rpm
    points = 2**math.ceil(math.log2(max(32, 4*effective_max*ratio)))
    total_revolutions = float(phase[-1])
    if total_revolutions < 2:
        raise ValueError('所选区间少于 2 转；请延长采集或缩小转速范围后再做阶次分析。')
    count = int(math.floor(total_revolutions*points))+1
    if points > 65536 or count > 4_000_000:
        raise ValueError('转速跨度或区间时长超过角域采样上限；请缩小分析区间后重试。')
    # The lowest angular sampling rate is points*minimum frot. The cutoff
    # lies below its Nyquist and above all retained orders at the highest RPM.
    # Zero-phase Butterworth filtering avoids a phase delay relative to RPM.
    cutoff = min(.4*points*(min_rpm/60), .45*sample_rate)
    centred = samples-np.mean(samples)
    filtered = sosfiltfilt(butter(10, cutoff, fs=sample_rate, output='sos'), centred)
    # Fourfold bandlimited interpolation reduces high-frequency linear-grid
    # amplitude error. resample_poly applies its own zero-phase low-pass FIR.
    interpolation_signal = resample_poly(filtered, 4, 1)
    interpolation_times = absolute_times[0]+np.arange(len(interpolation_signal))/(4*sample_rate)
    angles = np.arange(count, dtype=float)/points
    target_times = np.interp(angles, phase, absolute_times)
    angular = np.interp(target_times, interpolation_times, interpolation_signal)
    return angular, points, effective_max


def _peaks(orders, db, reference_frequency, mode):
    finite = np.isfinite(db)
    if not np.any(finite & (orders > .1)):
        return ()
    finite_db = db[finite & (orders > .1)]
    overall_peak = float(np.max(finite_db))
    # Whole-record median is deliberately conservative: random white-noise
    # fluctuations are not treated as mechanical lines. These gates describe
    # resolvable tonal evidence, not severity or a calibrated fault threshold.
    baseline = float(np.median(finite_db))
    if overall_peak-baseline < 18 or overall_peak < -240:
        return ()
    indices, _ = find_peaks(np.where(finite, db, -300.))
    result = []
    for index in indices:
        if orders[index] <= .1 or db[index] < overall_peak-40:
            continue
        lo, hi = max(0, index-50), min(len(db), index+51)
        neighbours = np.concatenate((db[lo:max(lo, index-3)], db[min(hi,index+4):hi]))
        neighbours = neighbours[np.isfinite(neighbours)]
        if len(neighbours) < 4:
            continue
        prominence = float(db[index]-np.median(neighbours))
        if prominence < 15:
            continue
        # Parabolic interpolation in dB reduces off-grid frequency bias;
        # level remains the actual measured FFT-bin peak, not an invented fit.
        delta = 0.
        if 0 < index < len(db)-1 and np.isfinite(db[index-1:index+2]).all():
            left, middle, right = db[index-1:index+2]
            denominator = left-2*middle+right
            if abs(denominator) > 1e-12:
                delta = float(np.clip(.5*(left-right)/denominator, -.5, .5))
        order = float(orders[index]+delta*(orders[1]-orders[0]))
        result.append(RotationPeak(order*reference_frequency, order, float(db[index]), prominence,
                                   'observed' if mode == 'manual' else 'reference_equivalent'))
    return tuple(sorted(result, key=lambda item: item.level_db, reverse=True)[:24])


def _markers(config, reference_rpm, minimum_rpm, maximum_rpm):
    definitions = [('1X', 1.), ('2X', 2.), ('3X', 3.), ('输出轴 1X', 1/config.gear_ratio)]
    if config.gear_teeth:
        teeth = float(config.gear_teeth)
        definitions.extend([('GMF', teeth), ('GMF−1X', teeth-1), ('GMF+1X', teeth+1)])
    if config.pole_pairs:
        definitions.append(('电磁 2pX', float(2*config.pole_pairs)))
    return tuple(RotationMarker(label, order*reference_rpm/60, order,
                                order*minimum_rpm/60, order*maximum_rpm/60)
                 for label, order in definitions if order > 0)


def _findings(peaks, config, effective_max, resolution, mode):
    if resolution > .125+1e-12:
        return ()
    tolerance = max(.035, 2*resolution)
    def match_window(order):
        # Adjacent ±1X mesh sidebands must never be accepted as the GMF itself.
        return max(tolerance, min(.25, .01*order))
    def matching(order):
        return next((p for p in peaks if abs(p.order-order) <= match_window(order)), None)
    def observation(p):
        suffix = '参考转速等效 Hz' if mode == 'measured' else 'Hz'
        return f'观测 {p.order:.3f}X，{p.frequency:.2f} {suffix}，峰值 {p.level_db:.2f} dB，邻域突出 {p.prominence_db:.1f} dB。'
    assumption = ('转速来自同步 CSV，使用角域重采样。' if mode == 'measured' else
                  '阶次按手动转速的定速假设换算；需用转速变化或转速计验证同步关系。')
    result = []
    one = matching(1.)
    if one:
        result.append(RotationFinding('转频同步激励候选（不平衡、偏心或安装传递）',
            ('1X 附近存在可分辨窄带峰。', observation(one), assumption),
            ('改变转速，检查该峰是否随 1X 移动；停机录制背景以排除 50/60 Hz 环境电声。',
             '检查电机轴/联轴器偏心、转子平衡和座椅安装点；用振动测量区分激励与传递路径。'),
            ('证实偏心或不平衡后校正同轴度/平衡；证实安装传递后优化紧固、刚度或隔振。',)))
    two = matching(2.)
    if two:
        result.append(RotationFinding('2X 谐波激励候选（对中、间隙或非线性接触）',
            ('2X 附近存在可分辨窄带峰。', observation(two), assumption),
            ('检查联轴器对中、轴承/导轨间隙及紧固件；对比正反转、空载与负载下的 2X。',
             '用运行转速变化确认该峰是同步谐波，而非固定频率背景声或结构共振。'),
            ('验证具体接触/对中问题后调整间隙、同轴度及紧固；复测相同负载和麦克风位置。',)))
    mesh = matching(float(config.gear_teeth)) if config.gear_teeth else None
    if mesh:
        lower, upper = matching(config.gear_teeth-1.), matching(config.gear_teeth+1.)
        evidence = ['提供齿数与输入轴转速推算的 GMF 附近存在峰。', observation(mesh), assumption]
        if lower and upper:
            evidence.extend(('GMF 两侧存在相隔约 1X 的成对侧带，提示啮合幅值/相位调制。', observation(lower), observation(upper)))
        result.append(RotationFinding('齿轮啮合调制候选' if lower and upper else '齿轮啮合激励候选',
            tuple(evidence),
            ('核对输入轴参考转速、实际啮合齿数/级数与传动比，避免把不同轴的 1X 当作啮合基准。',
             '改变载荷和转向，比较 GMF 与侧带；检查齿面接触、偏心、齿隙、润滑与装配公差。'),
            ('确认齿轮接触或装配问题后调整接触、齿隙、同轴度和润滑；用同工况录音/振动复核侧带变化。',)))
    electrical = matching(float(2*config.pole_pairs)) if config.pole_pairs else None
    if electrical:
        result.append(RotationFinding('电磁相关同步激励候选',
            ('用户提供极对数对应的 2pX 理论位置附近存在峰。', observation(electrical), assumption),
            ('同步测量驱动电流/电压，并核对电机类型、极对数、换向方式和控制频率。',
             '改变电源或控制参数并观察峰值；2pX 并非所有电机的唯一电磁噪声规律。'),
            ('确认电磁激励来源后针对驱动、换向或定转子间隙改善；保持机械工况一致进行复测。',)))
    known = (1., 2., float(config.gear_teeth), float(2*config.pole_pairs),
             float(config.gear_teeth-1) if config.gear_teeth else 0.,
             float(config.gear_teeth+1) if config.gear_teeth else 0.)
    high = next((p for p in peaks if p.frequency >= 500 and p.order >= 10
                 and all(abs(p.order-order) > match_window(order) for order in known if order)), None)
    if high:
        result.append(RotationFinding('窄带高频声或结构响应候选',
            (observation(high), '单个声音窄带峰不足以确认轴承损伤、啸叫或结构共振。'),
            ('做转速扫频：随转速移动提示同步激励，固定 Hz 增强提示背景或结构响应。',
             '移动麦克风、测量关键安装点振动，并与健康样件/停机背景对比定位源与传递路径。'),
            ('确认接触摩擦后改善接触/润滑；确认结构响应后优化局部刚度、阻尼或隔振，再复测。',)))
    return tuple(result)


def analyze_rotation(samples: Iterable[float], sample_rate: float, config: RotationConfig,
                     *, rpm_times=None, rpm_values=None, time_origin=0., source_unit='FS') -> RotationResult:
    """Analyze one channel; RPM CSV time must cover all selected sample timestamps.

    Frequency spectra use single-sided Hann-window peak amplitude, reference
    1 source unit. No Pa calibration, SPL, motor health threshold or ML confidence
    is inferred from an uncalibrated microphone. map_db is [order, time].
    """
    _validate_config(config)
    x = np.asarray(samples, dtype=float)
    if x.ndim != 1 or len(x) < 64 or not np.isfinite(x).all():
        raise ValueError('转频分析需要至少 64 个有限数值的一维音频样本。')
    if not np.isfinite(sample_rate) or sample_rate <= 0 or not np.isfinite(time_origin):
        raise ValueError('采样率必须为有限正数，时间起点必须为有限数值。')
    sample_rate = float(sample_rate)
    time_origin = float(time_origin)
    absolute_times = time_origin+np.arange(len(x))/sample_rate
    measured = rpm_times is not None or rpm_values is not None
    if measured and (rpm_times is None or rpm_values is None):
        raise ValueError('实测转速需要同时提供 rpm_times 和 rpm_values。')
    notes = ['幅度使用单边 Hann 窗峰值谱，参考 1 个源单位；这里的 dB 不是声压级 SPL。',
             '所有机理均为候选，须结合转速、负载、背景和振动/电流验证；不输出自动故障判定。']
    if str(source_unit).upper() == 'FS':
        notes.append('声卡 FS 未经 Pa 校准，仅可用于频率/阶次定位及同条件相对比较。')
    frequency_spectrum = compute_spectrum(x, sample_rate, reference=1.)
    points = 0
    if measured:
        curve_times, curve_rpm = _rpm_arrays(rpm_times, rpm_values)
        if curve_times[0] > absolute_times[0]+1e-10 or curve_times[-1] < absolute_times[-1]-1e-10:
            raise ValueError('实测转速曲线必须覆盖所选音频的全部样本时间；不允许外推。')
        inside = (curve_times > absolute_times[0]) & (curve_times < absolute_times[-1])
        plot_times = np.concatenate(([absolute_times[0]], curve_times[inside], [absolute_times[-1]]))
        plot_rpm = np.interp(plot_times, curve_times, curve_rpm)
        minimum, maximum = float(np.min(plot_rpm)), float(np.max(plot_rpm))
        phase = _phase_at(absolute_times, curve_times, curve_rpm)
        reference_rpm = float(phase[-1]/(absolute_times[-1]-absolute_times[0])*60)
        angular, points, effective_max = _angular_signal(x, sample_rate, absolute_times, phase,
                                                        minimum, maximum, config.max_order)
        order_spectrum = compute_spectrum(angular, points, reference=1.)
        orders, order_db = order_spectrum['frequency'], order_spectrum['db']
        map_orders, map_times, map_db = _map(angular, points, 8*points, 1.,
                                            phase=phase, original_times=absolute_times)
        mode = 'measured'
        notes.extend(('使用同步 CSV 的分段线性 rpm 精确积分，抗混叠滤波后作等角度重采样；未估计未知转速。',
                      '实测模式仅使用 CSV 转速；配置中的手动 rpm 不参与计算。',
                      '实测阶次峰与标线的单值 Hz 是参考转速下的等效频率，变速过程请同时看转速曲线与阶次图。',
                      f'角域每转 {points} 点，参考转速 {reference_rpm:.3f} rpm；输入转速范围 {minimum:.3f}–{maximum:.3f} rpm。'))
    else:
        minimum = maximum = reference_rpm = float(config.rpm)
        frot = config.rpm/60
        if (len(x)-1)/sample_rate*frot < 2:
            raise ValueError('所选区间少于 2 转；请延长采集或选择更长事件后再做阶次换算。')
        orders = frequency_spectrum['frequency']/frot
        order_db = frequency_spectrum['db']
        effective_max = min(float(config.max_order), sample_rate/(2*frot))
        map_orders, map_times, map_db = _map(x, sample_rate, round(sample_rate*8/frot), frot, time_origin)
        plot_times = np.array([absolute_times[0], absolute_times[-1]])
        plot_rpm = np.array([config.rpm, config.rpm])
        mode = 'manual'
        notes.append('手动 rpm 仅作定速频率换算，不是实测转速，未测量转速；变速录音应导入同步转速 CSV。')
    if effective_max < config.max_order-1e-9:
        notes.append(f'受原采样率和最高转速限制，最大阶次已截断为 {effective_max:.3f}X。')
    keep = orders <= effective_max+1e-10
    orders, order_db = orders[keep], order_db[keep]
    map_keep = map_orders <= effective_max+1e-10
    map_orders, map_db = map_orders[map_keep], map_db[map_keep, :]
    peaks = _peaks(orders, order_db, reference_rpm/60, mode)
    resolution = float(orders[1]-orders[0]) if len(orders) > 1 else effective_max
    if resolution > .125+1e-12:
        notes.append(f'全区间阶次分辨率 {resolution:.3f}X，转数不足以分辨相邻同步机理；保留图谱但不生成机理候选。')
    findings = _findings(peaks, config, effective_max, resolution, mode)
    if not findings:
        notes.append('未观察到足够支持上述候选机理的同步窄带证据；这不代表样件无故障。')
    return RotationResult(*map(_readonly, (frequency_spectrum['frequency'], frequency_spectrum['db'],
        orders, order_db, map_times, map_orders, map_db, plot_times, plot_rpm)),
        _markers(config, reference_rpm, minimum, maximum), peaks, findings, tuple(notes), config,
        str(source_unit), sample_rate, time_origin, len(x)/sample_rate, mode,
        f'峰值幅度 / dB re 1 {source_unit}', effective_max, points, reference_rpm)


def _export_payload(result, source_path, source_hash, config):
    config = config or result.config
    if config != result.config:
        raise ValueError('导出参数与已分析结果不一致；请重新分析后导出。')
    def serializable(value):
        if isinstance(value, np.ndarray):
            return serializable(value.tolist())
        if isinstance(value, np.generic):
            return serializable(value.item())
        if isinstance(value, (tuple, list)):
            return [serializable(item) for item in value]
        if isinstance(value, dict):
            return {key: serializable(item) for key, item in value.items()}
        if isinstance(value, (float, np.floating)) and not np.isfinite(value):
            return None
        return value
    return serializable({'schema_version': 1, 'source_path': str(source_path), 'source_hash': str(source_hash),
                         **{name: getattr(result, name) for name in (
                             'frequency', 'spectrum_db', 'orders', 'order_db', 'map_times', 'map_orders',
                             'map_db', 'rpm_times', 'rpm_values', 'notes', 'source_unit', 'sample_rate',
                             'time_origin', 'duration', 'rpm_mode', 'amplitude_label', 'effective_max_order',
                             'samples_per_revolution', 'frequency_reference_rpm')},
                         'config': asdict(config), 'markers': [asdict(m) for m in result.markers],
                         'peaks': [asdict(p) for p in result.peaks],
                         'findings': [asdict(f) for f in result.findings]})


def export_rotation_json(result: RotationResult, path: str | Path, *, source_path='', source_hash='', config=None):
    """Export actual arrays and parameters; JSON null represents −inf silence dB."""
    payload = _export_payload(result, source_path, source_hash, config)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2), encoding='utf-8')
    return path


def _figure_png(result):
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    figure = Figure(figsize=(10, 7), constrained_layout=True)
    FigureCanvasAgg(figure)
    frequency_axes, order_axes, map_axes, rpm_axes = figure.subplots(2, 2).flat
    frequency_axes.plot(result.frequency, result.spectrum_db, linewidth=.7)
    frequency_axes.set(title='Observed frequency spectrum', xlabel='Frequency / Hz', ylabel=f'Peak dB re 1 {result.source_unit}')
    frequency_axes.set_xlim(0, min(result.sample_rate/2, max(1000., result.effective_max_order*result.frequency_reference_rpm/60)))
    order_axes.plot(result.orders, result.order_db, linewidth=.7)
    order_axes.set(title='Order spectrum' if result.rpm_mode == 'measured' else 'Constant RPM frequency conversion',
                   xlabel='Order / X', ylabel=f'Peak dB re 1 {result.source_unit}', xlim=(0, result.effective_max_order))
    for marker in result.markers:
        if marker.order <= result.effective_max_order:
            order_axes.axvline(marker.order, linewidth=.5, color='gray', alpha=.6)
    finite = result.map_db[np.isfinite(result.map_db)]
    maximum = float(np.max(finite)) if len(finite) else 0.
    image = map_axes.pcolormesh(result.map_times, result.map_orders, result.map_db,
                               shading='auto', vmin=maximum-70, vmax=maximum, cmap='viridis')
    map_axes.set(title='Order-time map', xlabel='Audio time / s', ylabel='Order / X')
    figure.colorbar(image, ax=map_axes, label='Peak dB')
    rpm_axes.plot(result.rpm_times, result.rpm_values)
    rpm_axes.set(title='Synchronous measured RPM' if result.rpm_mode == 'measured' else 'Manual constant RPM assumption',
                 xlabel='Audio time / s', ylabel='RPM')
    for axes in (frequency_axes, order_axes, rpm_axes):
        axes.grid(alpha=.2)
    buffer = io.BytesIO()
    figure.savefig(buffer, format='png', dpi=140)
    return buffer.getvalue()


def export_rotation_report(result: RotationResult, path: str | Path, *, source_path='', source_hash='', config=None):
    """Export DOCX or standalone HTML containing observed plots and review steps."""
    payload = _export_payload(result, source_path, source_hash, config)
    path = Path(path)
    if path.suffix.lower() not in ('.docx', '.html', '.htm'):
        raise ValueError('现场诊断报告请选择 .docx 或 .html。')
    image = _figure_png(result)
    metadata = [f'源文件：{source_path or "未提供"}', f'SHA-256：{source_hash or "未提供"}',
                f'区间：{result.time_origin:.6f}–{result.time_origin+result.duration:.6f} s；采样率：{result.sample_rate:g} Hz',
                f'转速模式：{"同步 CSV 实测" if result.rpm_mode == "measured" else "手动定速假设"}；参考转速：{result.frequency_reference_rpm:.3f} rpm',
                '原始参数：'+json.dumps(payload['config'], ensure_ascii=False), result.amplitude_label]
    if path.suffix.lower() in ('.html', '.htm'):
        esc = html.escape
        sections = ['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>现场声学与旋转机械证据报告</title>',
                    '<style>body{font-family:Microsoft YaHei,sans-serif;max-width:1000px;margin:32px auto;line-height:1.7}img{width:100%}li{margin:.4em 0}</style>',
                    '<h1>现场声学与旋转机械证据报告</h1>', *[f'<p>{esc(line)}</p>' for line in metadata],
                    '<img alt="实测频谱、阶次、阶次图与转速" src="data:image/png;base64,'+base64.b64encode(image).decode('ascii')+'">',
                    '<h2>观测峰</h2><ul>', *[f'<li>{p.order:.4f}X / {p.frequency:.3f} Hz / {p.level_db:.2f} dB / 邻域突出 {p.prominence_db:.1f} dB ({esc(p.frequency_kind)})</li>' for p in result.peaks], '</ul>']
        for finding in result.findings:
            sections.append(f'<h2>{esc(finding.mechanism)}</h2>')
            for label, entries in (('证据', finding.evidence), ('验证办法', finding.checks), ('改善建议', finding.recommendations)):
                sections.extend([f'<h3>{label}</h3><ul>', *[f'<li>{esc(entry)}</li>' for entry in entries], '</ul>'])
        sections.extend(['<h2>解释与限制</h2><ul>', *[f'<li>{esc(note)}</li>' for note in result.notes], '</ul></html>'])
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text('\n'.join(sections), encoding='utf-8')
    else:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
        document = Document()
        style = document.styles['Normal']
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(10)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        document.add_heading('现场声学与旋转机械证据报告', 0)
        for line in metadata:
            document.add_paragraph(line)
        document.add_picture(io.BytesIO(image), width=Inches(6.2))
        document.add_heading('观测峰', 1)
        table = document.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        for cell, title in zip(table.rows[0].cells, ('阶次 / X', '频率 / Hz', '峰值 / dB', '邻域突出 / dB', '频率含义')):
            cell.text = title
        for peak in result.peaks:
            for cell, value in zip(table.add_row().cells, (f'{peak.order:.4f}', f'{peak.frequency:.3f}',
                f'{peak.level_db:.2f}', f'{peak.prominence_db:.1f}', '观测定速 Hz' if peak.frequency_kind == 'observed' else '参考转速等效 Hz')):
                cell.text = value
        if not result.findings:
            document.add_paragraph('没有足够同步窄带证据支持具体候选机理；不代表样件无故障。')
        for finding in result.findings:
            document.add_heading(finding.mechanism, 1)
            for label, entries in (('证据', finding.evidence), ('验证办法', finding.checks), ('改善建议', finding.recommendations)):
                document.add_heading(label, 2)
                for entry in entries:
                    document.add_paragraph(entry, style='List Bullet')
        document.add_heading('解释与限制', 1)
        for note in result.notes:
            document.add_paragraph(note, style='List Bullet')
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)
    return path
