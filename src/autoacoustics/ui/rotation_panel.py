"""Explicit RPM/order diagnostics with immutable input snapshots and stale-result guards."""
from __future__ import annotations

from pathlib import Path
import hashlib
import json
import math

import numpy as np
import pyqtgraph as pg
from PyQt6.QtCore import QThread, QRectF, pyqtSignal
from PyQt6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
    QGroupBox, QLabel, QComboBox, QLineEdit, QPushButton, QFormLayout,
    QDoubleSpinBox, QSpinBox, QFileDialog, QTableWidget, QTableWidgetItem,
    QHeaderView, QDialog, QScrollArea,QTabWidget,QSizePolicy)

from autoacoustics.analysis.rotation import (RotationConfig, analyze_rotation, load_rpm_csv,
    export_rotation_json, export_rotation_report)
from .workbench import WORKBENCH_STYLE


class _RotationTask(QThread):
    succeeded = pyqtSignal(object)
    failed = pyqtSignal(str)

    def __init__(self, function, parent=None):
        super().__init__(parent)
        self.function = function

    def run(self):
        try:
            self.succeeded.emit(self.function())
        except Exception as error:
            self.failed.emit(str(error))


class RotationPanel(QWidget):
    overlayReady = pyqtSignal(object)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName('rotationPanel')
        self.setStyleSheet(WORKBENCH_STYLE)
        self.signal = self.result = None
        self.channel = 0
        self.selection_seconds = None
        self.result_stale = True
        self._generation = 0
        self._source_key = None
        self._worker = self._export_worker = None
        self._closing = False
        self._result_snapshot = None
        self._plot_data = {}
        self._dialogs = []
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        heading = QLabel('转频 / 阶次 · 有证据的候选机理')
        heading.setProperty('role', 'sectionTitle')
        layout.addWidget(heading)
        self.source_label = QLabel('尚未选择信号；分析沿用主窗口的原采样率、通道和事件区间。')
        self.source_label.setWordWrap(True)
        self.source_label.setMaximumHeight(42)
        self.source_label.setSizePolicy(QSizePolicy.Policy.Ignored,QSizePolicy.Policy.Preferred)
        layout.addWidget(self.source_label)
        grid = QGridLayout()
        layout.addLayout(grid, 1)
        self.order_plot = self._plot_card(grid, 0, 0, '阶次谱', 'orders')
        self.map_plot = self._plot_card(grid, 0, 1, '阶次历程 / 图', 'map')
        self.rpm_plot = self._plot_card(grid, 1, 0, '转速曲线 / RPM', 'rpm')
        self.order_plot.setLabel('bottom', '阶次', units='order')
        self.order_plot.setLabel('left', '幅度')
        self.map_plot.setLabel('bottom', '原录音时间', units='s')
        self.map_plot.setLabel('left', '阶次', units='order')
        self.rpm_plot.setLabel('bottom', '原录音时间', units='s')
        self.rpm_plot.setLabel('left', '转速', units='rpm')
        self.order_curve = self.order_plot.plot(pen=pg.mkPen('#67d7e3', width=1))
        self.rpm_curve = self.rpm_plot.plot(pen=pg.mkPen('#e2c25e', width=1))
        self.image = pg.ImageItem(axisOrder='row-major')
        self.map_plot.addItem(self.image)
        self.image.setLookupTable(pg.colormap.get('viridis').getLookupTable())
        control = QGroupBox('明确参数与转速来源')
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(control)
        self.controls_tabs=QTabWidget()
        self.controls_tabs.addTab(scroll,'转速 / 传动参数')
        results_page=QWidget();results_layout=QVBoxLayout(results_page)
        results_layout.setContentsMargins(4,4,4,4)
        self.results_scroll=QScrollArea()
        self.results_scroll.setWidgetResizable(True)
        self.results_scroll.setWidget(results_page)
        self.controls_tabs.addTab(self.results_scroll,'诊断证据 / 建议')
        grid.addWidget(self.controls_tabs,1,1)
        form = QFormLayout(control)
        form.setRowWrapPolicy(QFormLayout.RowWrapPolicy.WrapLongRows)
        self.rpm_source_combo = QComboBox()
        self.rpm_source_combo.addItem('手动恒速 RPM（假设）', 'manual')
        self.rpm_source_combo.addItem('实测 RPM CSV（原时间轴）', 'measured')
        self.rpm_spin = QDoubleSpinBox()
        self.rpm_spin.setRange(.001, 1000000)
        self.rpm_spin.setDecimals(3)
        self.rpm_spin.setValue(3000)
        self.rpm_spin.setSuffix(' rpm')
        self.rpm_csv_path = QLineEdit()
        self.rpm_csv_path.setPlaceholderText('time_s,rpm；须覆盖所选原时间区间')
        csv_button = QPushButton('选择 CSV')
        csv_button.clicked.connect(self._choose_csv)
        row = QHBoxLayout()
        row.addWidget(self.rpm_csv_path, 1)
        row.addWidget(csv_button)
        self.gear_teeth_spin, self.pole_pairs_spin = QSpinBox(), QSpinBox()
        self.gear_teeth_spin.setRange(0, 10000)
        self.pole_pairs_spin.setRange(0, 1000)
        self.gear_teeth_spin.setSpecialValueText('未知 / 不计算 GMF')
        self.pole_pairs_spin.setSpecialValueText('未知 / 不计算电磁阶次')
        self.gear_ratio_spin, self.max_order_spin = QDoubleSpinBox(), QDoubleSpinBox()
        self.gear_ratio_spin.setRange(.0001, 1000000)
        self.gear_ratio_spin.setDecimals(4)
        self.gear_ratio_spin.setValue(1)
        self.gear_ratio_spin.setToolTip('输入轴转速 / 被观测轴转速；本面板按分析引擎定义记录。')
        self.max_order_spin.setRange(.1, 10000)
        self.max_order_spin.setValue(50)
        form.addRow('转速来源', self.rpm_source_combo)
        form.addRow('手动 RPM', self.rpm_spin)
        form.addRow('实测转速文件', row)
        form.addRow('齿数', self.gear_teeth_spin)
        form.addRow('传动比', self.gear_ratio_spin)
        form.addRow('极对数', self.pole_pairs_spin)
        form.addRow('最大分析阶次', self.max_order_spin)
        notice = QLabel('手动 RPM 只支持恒速假设；实测 CSV 做变速阶次跟踪。峰值匹配只是候选证据，不自动判故障或 OK/NG。FS 未校准时只保留相对幅度。')
        notice.setWordWrap(True)
        form.addRow(notice)
        self.analyze_button = QPushButton('分析当前通道 / 选段')
        self.analyze_button.setProperty('role', 'primary')
        self.analyze_button.clicked.connect(self.start_analysis)
        self.export_button = QPushButton('导出诊断报告 / 原参数 JSON')
        self.export_button.clicked.connect(self._choose_export)
        form.addRow(self.analyze_button)
        form.addRow(self.export_button)
        self.findings_table = QTableWidget(0, 4)
        self.findings_table.setHorizontalHeaderLabels(['候选机理', '证据', '验证步骤', '建议措施'])
        self.findings_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.findings_table.setWordWrap(True)
        self.findings_table.setMinimumHeight(150)
        results_layout.addWidget(self.findings_table,1)
        self.peaks_label = QLabel('峰值候选：等待分析。')
        self.peaks_label.setWordWrap(True)
        self.peaks_label.setSizePolicy(QSizePolicy.Policy.Ignored,QSizePolicy.Policy.Preferred)
        results_layout.addWidget(self.peaks_label)
        self.status_label = QLabel('等待信号；未知齿数和极对数不推断。')
        self.status_label.setWordWrap(True)
        self.status_label.setSizePolicy(QSizePolicy.Policy.Ignored,QSizePolicy.Policy.Preferred)
        results_layout.addWidget(self.status_label)
        result_actions=QHBoxLayout()
        repeat=QPushButton('返回参数');repeat.clicked.connect(lambda:self.controls_tabs.setCurrentIndex(0))
        export=QPushButton('导出诊断报告');export.clicked.connect(self._choose_export)
        result_actions.addWidget(repeat);result_actions.addWidget(export);results_layout.addLayout(result_actions)
        for widget in (self.rpm_spin, self.gear_teeth_spin, self.gear_ratio_spin,
                       self.pole_pairs_spin, self.max_order_spin):
            widget.valueChanged.connect(self._parameters_changed)
        self.rpm_source_combo.currentIndexChanged.connect(self._parameters_changed)
        self.rpm_csv_path.textChanged.connect(self._parameters_changed)
        self._enable_controls()
        grid.setColumnStretch(0, 1)
        grid.setColumnStretch(1, 1)
        grid.setRowStretch(0, 1)
        grid.setRowStretch(1, 1)

    def _plot_card(self, grid, row, column, title, key):
        group = QGroupBox(title)
        layout = QVBoxLayout(group)
        plot = pg.PlotWidget(background='#101018')
        plot.setMinimumSize(180, 105)
        layout.addWidget(plot, 1)
        expand = QPushButton('放大图表')
        expand.clicked.connect(lambda:self._expand(key, title))
        layout.addWidget(expand)
        grid.addWidget(group, row, column)
        return plot

    def _expand(self, key, title):
        dialog = QDialog(self)
        dialog.setWindowTitle(title+(' · 结果已过期' if self.result_stale else ''))
        dialog.setStyleSheet(WORKBENCH_STYLE)
        dialog.resize(960, 620)
        layout = QVBoxLayout(dialog)
        plot = pg.PlotWidget(background='#101018')
        layout.addWidget(plot)
        value = self._plot_data.get(key)
        if value is not None:
            if key == 'map':
                data, rect, levels = value
                image = pg.ImageItem(data, axisOrder='row-major')
                image.setLookupTable(pg.colormap.get('viridis').getLookupTable())
                image.setRect(rect)
                image.setLevels(levels)
                plot.addItem(image)
            else:
                plot.plot(value[0], value[1], pen='#67d7e3')
        self._dialogs.append(dialog)
        dialog.finished.connect(lambda:self._dialogs.remove(dialog) if dialog in self._dialogs else None)
        dialog.show()

    @property
    def is_analyzing(self):
        return self._worker is not None

    def _enable_controls(self):
        self.analyze_button.setEnabled(self.signal is not None and not self.is_analyzing and not self._closing)
        self.export_button.setEnabled(bool(self.result is not None and not self.result_stale and
                                           not self.is_analyzing and self._export_worker is None))
        measured = self.rpm_source_combo.currentData() == 'measured'
        self.rpm_spin.setEnabled(not measured)
        self.rpm_csv_path.setEnabled(measured)

    def _parameters_changed(self, *unused):
        self._invalidate('分析参数或转速来源已变化；旧结果已过期，请重新分析。')
        self._enable_controls()

    def _invalidate(self, reason, *, clear=False):
        self._generation += 1
        self.result_stale = True
        self.status_label.setText(reason)
        self.export_button.setEnabled(False)
        self.overlayReady.emit({'stale':True, 'source_hash':self.signal.source_hash if self.signal else '',
                                'channel':self.channel, 'selection_seconds':self.selection_seconds})
        if clear:
            self.result = None
            self._result_snapshot = None
            self.order_curve.clear()
            self.rpm_curve.clear()
            self.image.clear()
            self.findings_table.setRowCount(0)
            self._plot_data.clear()

    def set_signal(self, signal, channel=0, start_seconds=None, end_seconds=None):
        if signal is None:
            self.signal = None
            self._source_key = None
            self.selection_seconds = None
            self.source_label.setText('尚未选择信号。')
            self._invalidate('没有信号，旧诊断已清除。', clear=True)
            self._enable_controls()
            return
        if not isinstance(channel, int) or not 0 <= channel < signal.channel_count:
            raise ValueError('转频诊断通道索引无效。')
        origin = signal.time_origin
        start_seconds = origin if start_seconds is None else float(start_seconds)
        end_seconds = origin+signal.frames/signal.sample_rate if end_seconds is None else float(end_seconds)
        start = round((start_seconds-origin)*signal.sample_rate)
        end = round((end_seconds-origin)*signal.sample_rate)
        if not all(math.isfinite(v) for v in (start_seconds, end_seconds)) or not 0 <= start < end <= signal.frames:
            raise ValueError('转频选段必须在原录音时间轴内且至少有一个样本。')
        selection = (origin+start/signal.sample_rate, origin+end/signal.sample_rate)
        key = (id(signal), channel, start, end, signal.source_hash, signal.sample_rate)
        if key == self._source_key:
            return
        self.signal, self.channel = signal, channel
        self.selection_seconds = selection
        self._selected_samples = signal.samples[channel, start:end]
        self._source_key = key
        self.source_label.setText(f'来源 {signal.path or "未存盘"} · SHA256 {signal.source_hash or "未提供"}\n'+
            f'通道 {channel}: {signal.channels[channel].name} / {signal.channels[channel].unit} · '+
            f'{signal.sample_rate:g} Hz · 原时间 [{selection[0]:g}, {selection[1]:g}) s · 原样本 [{start}, {end})')
        self._invalidate('来源、通道或选段已变化，请分析当前输入。', clear=True)
        self._enable_controls()

    def _choose_csv(self):
        path, _ = QFileDialog.getOpenFileName(self, '选择实测转速 CSV', '', 'RPM CSV (*.csv);;所有文件 (*)')
        if path:
            self.rpm_csv_path.setText(path)
            self.rpm_source_combo.setCurrentIndex(self.rpm_source_combo.findData('measured'))

    def start_analysis(self):
        if self.signal is None or self.is_analyzing or self._closing:
            return
        config = RotationConfig(rpm=self.rpm_spin.value(), gear_teeth=self.gear_teeth_spin.value(),
            gear_ratio=self.gear_ratio_spin.value(), pole_pairs=self.pole_pairs_spin.value(),
            max_order=self.max_order_spin.value())
        generation = self._generation
        mode = self.rpm_source_combo.currentData()
        csv_path = self.rpm_csv_path.text().strip() if mode == 'measured' else ''
        if mode == 'measured' and not csv_path:
            self.status_label.setText('请明确选择 time_s,rpm CSV，时间须覆盖所选原时间区间。')
            return
        snapshot = {'generation':generation, 'source_path':str(self.signal.path or ''),
            'source_hash':self.signal.source_hash, 'channel':self.channel,
            'selection_seconds':self.selection_seconds, 'source_unit':self.signal.channels[self.channel].unit,
            'sample_rate':self.signal.sample_rate, 'config':config, 'rpm_mode':mode,
            'rpm_csv_path':csv_path}
        samples = self._selected_samples
        def analyze():
            rpm_times = rpm_values = None
            if csv_path:
                path = Path(csv_path)
                snapshot['rpm_csv_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
                rpm_times, rpm_values = load_rpm_csv(path)
                if hashlib.sha256(path.read_bytes()).hexdigest()!=snapshot['rpm_csv_sha256']:
                    raise ValueError('读取期间转速 CSV 已改变，请确认文件稳定后重新分析。')
            result = analyze_rotation(samples, snapshot['sample_rate'], config,
                rpm_times=rpm_times, rpm_values=rpm_values, time_origin=snapshot['selection_seconds'][0],
                source_unit=snapshot['source_unit'])
            return result, snapshot
        self.result_stale = True
        self.overlayReady.emit({'stale':True, 'source_hash':snapshot['source_hash'],
                                'channel':snapshot['channel'], 'selection_seconds':snapshot['selection_seconds']})
        self.status_label.setText('后台分析中；保持原采样率与选段，参数变化将使本次结果失效。')
        self._worker = _RotationTask(analyze, self)
        self._worker.succeeded.connect(self._analysis_completed)
        self._worker.failed.connect(self._analysis_failed)
        self._worker.finished.connect(self._worker_finished)
        self._worker.start()
        self._enable_controls()

    def _analysis_completed(self, payload):
        result, snapshot = payload
        if snapshot['generation'] != self._generation:
            self.status_label.setText('后台结果对应旧来源 / 参数，已丢弃；请分析当前输入。')
            return
        self.result, self._result_snapshot = result, snapshot
        self.result_stale = False
        self._show_result(result)
        markers = [{'label':marker.label, 'frequency':marker.frequency, 'order':marker.order,
                    'frequency_low':getattr(marker, 'frequency_low', marker.frequency),
                    'frequency_high':getattr(marker, 'frequency_high', marker.frequency)} for marker in result.markers]
        self.overlayReady.emit({'stale':False, 'source_hash':snapshot['source_hash'],
            'channel':snapshot['channel'], 'selection_seconds':snapshot['selection_seconds'],
            'rpm_mode':result.rpm_mode, 'time_origin':result.time_origin, 'markers':markers,
            'rpm_times':result.rpm_times, 'rpm_values':result.rpm_values,
            'frequency':result.frequency, 'spectrum_db':result.spectrum_db, 'result':result})
        self.status_label.setText('诊断完成；'+('实测变速阶次' if result.rpm_mode == 'measured' else '手动恒速假设')+
            '。候选不代表已确认故障。\n'+'；'.join(result.notes))
        self.status_label.setToolTip(self.status_label.text())
        self.controls_tabs.setCurrentIndex(1)

    def _analysis_failed(self, message):
        self.status_label.setText(f'诊断未完成：{message}；请核对 RPM、CSV 时间覆盖与选段长度。')
        self.result_stale = True

    def _worker_finished(self):
        if self._worker:
            self._worker.deleteLater()
            self._worker = None
        self._enable_controls()

    @staticmethod
    def _lines(value):
        return value if isinstance(value, str) else '\n'.join(map(str, value))

    def _show_result(self, result):
        self.order_plot.setLabel('left', result.amplitude_label)
        self.order_curve.setData(result.orders, result.order_db)
        self._plot_data['orders'] = (result.orders, result.order_db)
        self.rpm_curve.setData(result.rpm_times, result.rpm_values)
        self._plot_data['rpm'] = (result.rpm_times, result.rpm_values)
        if result.map_db.size and len(result.map_times) and len(result.map_orders):
            display_times=np.linspace(result.map_times[0],result.map_times[-1],len(result.map_times))
            display_db=result.map_db
            interpolated=len(display_times)>1 and not np.allclose(display_times,result.map_times,rtol=1e-8,atol=1e-10)
            if interpolated:
                display_db=np.array([np.interp(display_times,result.map_times,row) for row in result.map_db])
            self.map_plot.setLabel('bottom','原录音时间 · 显示时间插值' if interpolated else '原录音时间',units='s')
            dt = float(display_times[1]-display_times[0]) if len(display_times) > 1 else result.duration
            do = float(np.median(np.diff(result.map_orders))) if len(result.map_orders) > 1 else 1.
            rect = QRectF(float(result.map_times[0]-dt/2), float(result.map_orders[0]-do/2),
                          max(dt, len(result.map_times)*dt), max(do, len(result.map_orders)*do))
            finite = result.map_db[np.isfinite(result.map_db)]
            high = float(np.max(finite)) if finite.size else 0.
            levels = (high-80, high)
            self.image.setImage(display_db, autoLevels=False)
            self.image.setRect(rect)
            self.image.setLevels(levels)
            self._plot_data['map'] = (display_db, rect, levels)
            self.map_plot.setXRange(rect.left(), rect.right(), padding=0)
            self.map_plot.setYRange(max(0, rect.top()), rect.bottom(), padding=0)
        self.findings_table.setRowCount(len(result.findings))
        for row, finding in enumerate(result.findings):
            for column, value in enumerate((finding.mechanism, finding.evidence, finding.checks, finding.recommendations)):
                item = QTableWidgetItem(self._lines(value))
                item.setToolTip(self._lines(value))
                self.findings_table.setItem(row, column, item)
        self.findings_table.resizeRowsToContents()
        self.peaks_label.setText('峰值候选（'+('参考转速等效 Hz' if result.rpm_mode == 'measured' else '实测频率 Hz')+'）：'+
            '；'.join(f'{peak.order:.3g} 阶 / {peak.frequency:.3g} Hz / {peak.level_db:.1f} dB' for peak in result.peaks[:10]))

    def _export_snapshot(self):
        if self.result is None or self.result_stale or self.is_analyzing or self._result_snapshot is None:
            raise ValueError('诊断结果已过期或尚未完成，请重新分析当前来源和参数。')
        return self.result, dict(self._result_snapshot)

    @staticmethod
    def _write_export(result, snapshot, path):
        path = Path(path)
        args = {'source_path':snapshot['source_path'], 'source_hash':snapshot['source_hash'],
                'config':snapshot['config']}
        if path.suffix.lower() == '.json':
            saved = export_rotation_json(result, path, **args)
            data = json.loads(Path(saved).read_text(encoding='utf-8'))
            data['ui_source'] = {name:snapshot.get(name) for name in
                ('source_path', 'source_hash', 'channel', 'selection_seconds', 'source_unit',
                 'sample_rate', 'rpm_mode', 'rpm_csv_path', 'rpm_csv_sha256')}
            Path(saved).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
            return Path(saved)
        if path.suffix.lower() not in {'.docx', '.html'}:
            raise ValueError('诊断导出请选择 .docx、.html 或 .json。')
        saved=Path(export_rotation_report(result,path,**args))
        details=[f'通道索引：{snapshot["channel"]}',f'源单位：{snapshot["source_unit"]}',
            f'原选段秒区间：{snapshot["selection_seconds"]}',f'转速 CSV：{snapshot.get("rpm_csv_path") or "无；手动恒速假设"}',
            f'转速 CSV SHA-256：{snapshot.get("rpm_csv_sha256") or "不适用"}']
        if saved.suffix.lower()=='.docx':
            from docx import Document
            document=Document(saved);document.add_heading('采集与转速来源快照',level=1)
            for line in details:document.add_paragraph(line)
            document.save(saved)
        else:
            import html
            content=saved.read_text(encoding='utf-8')
            block='<h2>采集与转速来源快照</h2>'+''.join('<p>'+html.escape(line)+'</p>' for line in details)
            saved.write_text(content.replace('</html>',block+'</html>'),encoding='utf-8')
        return saved

    def export_result(self, path):
        result, snapshot = self._export_snapshot()
        return self._write_export(result, snapshot, path)

    def _choose_export(self):
        try:
            result, snapshot = self._export_snapshot()
        except ValueError as error:
            self.status_label.setText(str(error))
            return
        path, _ = QFileDialog.getSaveFileName(self, '导出诊断与原始参数', 'rotation.docx',
            'Word 报告 (*.docx);;原参数与结果 JSON (*.json);;HTML 报告 (*.html)')
        if not path:
            return
        self._export_worker = _RotationTask(lambda:self._write_export(result, snapshot, path), self)
        self._export_worker.succeeded.connect(self._export_completed)
        self._export_worker.failed.connect(self._export_failed)
        self._export_worker.finished.connect(self._export_finished)
        self._export_worker.start()
        self._enable_controls()

    def _export_completed(self, path):
        self.status_label.setText(f'已导出分析时的原参数 / 来源快照：{path}'+
                                  ('；当前界面已变化，此文件保留原分析上下文。' if self.result_stale else ''))

    def _export_failed(self, message):
        self.status_label.setText(f'报告导出失败：{message}；请检查目标目录和文件权限。')

    def _export_finished(self):
        if self._export_worker:
            self._export_worker.deleteLater()
            self._export_worker = None
        self._enable_controls()

    def shutdown(self):
        self._closing = True
        if self.is_analyzing or self._export_worker is not None:
            return False
        for dialog in tuple(self._dialogs):
            dialog.close()
        return True
