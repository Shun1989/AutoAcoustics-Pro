from pathlib import Path
import zipfile

from docx import Document
import numpy as np
import pytest

from autoacoustics.model import (AnalysisDetail, AnalysisResult, AnalysisSettings,
    AnalysisSummary, BatchItemResult, BatchResult, CancellationToken, ComparisonResult,
    MeasurementContext, MetricResult)
from autoacoustics import report


def result(detail=True, identity='left', value=81.23456):
    arrays = {'waveform_time': np.arange(200) / 48000,
              'waveform': np.sin(np.arange(200)),
              'frequency': np.arange(100), 'spectrum_db': np.arange(100) * 0.2,
              'psd': np.arange(100) * 0.001,
              'spl_time': np.arange(100) * .01, 'laf': np.arange(100) * .2 + 70,
              'las': np.arange(100) * .1 + 70,
              'loudness_time': np.arange(100) * .002, 'loudness': np.arange(100) * .01,
              'octave_centers': np.array([100., 1000., 10000.]),
              'octave_levels': np.array([55., 60., 57.]),
              'stft_time': np.arange(10) * .01,
              'stft_frequency': np.arange(15) * 100.,
              'stft_db': np.arange(150).reshape(15, 10) * .1}
    return AnalysisResult(identity, 'a' * 64, Path('座椅 中文.wav'),
        AnalysisSettings(channel=0, start=0, end=200),
        MeasurementContext(specimen='座椅电机', supply='12 V', repeat='第1次'), None,
        AnalysisSummary({'LZeq': MetricResult(value, 'dB', method='energy mean'),
            'LAeq': MetricResult(None, 'dB', 'uncalibrated', message='缺少校准'),
            'Nmax': MetricResult(0, 'sone', 'failed', message='计算失败')}),
        AnalysisDetail(arrays, {'waveform':'Pa', 'spectrum_db':'dB re 20 µPa',
            'psd':'Pa²/Hz', 'laf':'dB', 'las':'dB', 'loudness':'sone'}) if detail else None,
        {'sample_rate':48000,'source_unit':'Pa','unit':'Pa','algorithm_version':'0.1.0',
         'calibration':{'source':'文件声明 Pa'},'project_version':1,'event_end_sample':200})


def text(path):
    document = Document(path)
    return '\n'.join([p.text for p in document.paragraphs] +
                     [cell.text for table in document.tables for row in table.rows for cell in row.cells])


def test_single_report_uses_existing_results_states_configuration_and_images(tmp_path):
    path = report.write_docx(result(), tmp_path / '中文 单文件.docx')
    contents = text(path)
    assert '81.23456' in contents and 'dB' in contents
    assert '缺少校准' in contents and '计算失败' in contents
    assert '12 V' in contents and 'unknown' not in contents
    assert '[0, 200)' in contents and '48000' in contents
    assert '文件声明 Pa' in contents and '0.1.0' in contents
    assert '0 sone' not in contents
    with zipfile.ZipFile(path) as archive:
        assert len([name for name in archive.namelist() if name.startswith('word/media/')]) >= 6


def test_missing_detail_is_explicit_without_invented_curves(tmp_path):
    path = report.write_docx(result(False), tmp_path / '摘要.docx')
    assert '仅有摘要' in text(path)
    with zipfile.ZipFile(path) as archive:
        assert not any(name.startswith('word/media/') for name in archive.namelist())


def test_batch_has_all_rows_and_summary_chart_and_failures(tmp_path):
    batch = BatchResult((BatchItemResult('a', Path('中文.wav'), 'ok', result(False)),
                         BatchItemResult('b', Path('坏文件.wav'), 'failed', error='损坏文件')))
    path = report.write_docx(batch, tmp_path / '批量.docx')
    contents = text(path)
    assert '损坏文件' in contents and '坏文件.wav' in contents
    assert '81.23456' in contents and '仅有摘要' in contents
    with zipfile.ZipFile(path) as archive:
        assert any(name.startswith('word/media/') for name in archive.namelist())


def test_partial_batch_keeps_valid_metrics_in_charts_and_reports_failure_count(tmp_path):
    batch=BatchResult((BatchItemResult('partial-item',Path('部分指标.wav'),'partial',
          result(False),error='部分指标失败：Nmax'),))
    path=report.write_docx(batch,tmp_path/'部分指标.docx')
    contents=text(path)
    assert '部分指标失败 1 条' in contents and '状态 部分指标失败' in contents
    assert '完成 0 条' in contents and '81.23456' in contents
    assert '部分指标失败：Nmax' in contents and '计算失败' in contents
    with zipfile.ZipFile(path) as archive:
        # Its valid LZeq remains visible; uncalibrated LAeq and failed Nmax
        # must not become aggregate chart points or fake zero measurements.
        assert len([name for name in archive.namelist() if name.startswith('word/media/')])==1


def test_comparison_report_direction_units_conditions_and_roles(tmp_path):
    comparison = ComparisonResult(result(identity='left'), result(identity='right', value=84.23456),
          {'LZeq': MetricResult(3., 'dB', method='right minus left')},
          ('供电不同 左12 V 右14 V', '负载未知'))
    path = report.write_docx(comparison, tmp_path / '比较.docx')
    contents = text(path)
    assert '右侧减左侧' in contents and '供电不同' in contents and '负载未知' in contents
    assert 'left' in contents and 'right' in contents
    assert '3' in contents and 'dB' in contents


def test_atomic_failure_and_cancellation_preserve_previous_report(tmp_path, monkeypatch):
    target = tmp_path / '既有.docx'; target.write_bytes(b'previous report')
    token = CancellationToken(); token.cancel()
    with pytest.raises(Exception, match='取消'):
        report.write_docx(report.ReportSpec(result(False), cancel=token), target)
    assert target.read_bytes() == b'previous report'

    def fail(*args, **kwargs):
        raise OSError('disk failure')
    monkeypatch.setattr(report, '_save_document', fail)
    with pytest.raises(OSError):
        report.write_docx(result(False), target)
    assert target.read_bytes() == b'previous report'
    assert not list(tmp_path.glob('*.tmp.docx'))


def test_resolved_event_end_and_disabled_values_use_actual_pipeline_contract(tmp_path):
    from dataclasses import replace
    sample=replace(result(False),settings=AnalysisSettings(channel=0,end=None),
                   summary=AnalysisSummary({'Nstationary':MetricResult(0,'sone','disabled')}))
    path=report.write_docx(sample,tmp_path/'末尾.docx')
    contents=text(path)
    assert '[0, 200)' in contents
    assert '未启用' in contents and '0 sone' not in contents
    assert '文件声明 Pa' in contents
    assert '未校准或来源未记录' not in contents
    row=next(row for table in Document(path).tables for row in table.rows
             if 'Nstationary' in row.cells[0].text)
    assert row.cells[1].text=='—' and row.cells[3].text=='未启用'


def test_large_calibration_manifest_has_count_digest_and_current_input_identity(tmp_path):
    from dataclasses import replace
    import hashlib
    import json
    hashes=[hashlib.sha256(str(index).encode()).hexdigest() for index in range(64)]
    sample=replace(result(False),input_hash=hashes[0],provenance={
        'sample_rate':48000,'source_unit':'FS','unit':'Pa',
        'calibration':{'source':'本 ZIP 参考','applicable_hashes':hashes}})
    path=report.write_docx(sample,tmp_path/'大校准清单.docx')
    contents=text(path)
    digest=hashlib.sha256(json.dumps(hashes,ensure_ascii=False,sort_keys=True,
                                     separators=(',',':')).encode('utf-8')).hexdigest()
    assert hashes[0] in contents and digest in contents and '64 项' in contents
    assert hashes[-1] not in contents


def test_nonzero_source_origin_is_retained_in_report_times(tmp_path):
    from dataclasses import replace
    sample=replace(result(False),provenance={'sample_rate':48000,
        'time_origin_seconds':10.,'event_end_sample':200,
        'event_start_time_seconds':10.,'event_end_time_seconds':10.+200/48000})
    path=report.write_docx(sample,tmp_path/'源时钟.docx')
    assert '[10, 10.0041667) s' in text(path)


def test_batch_prints_shared_context_once_and_keeps_individual_summaries(tmp_path):
    from dataclasses import replace
    first=replace(result(False),context=MeasurementContext(specimen='共同唯一试件',supply='12 V'))
    second=replace(first,result_id='right',settings=replace(first.settings,start=10,event_label='第二事件'))
    batch=BatchResult((BatchItemResult('a',Path('同文件.wav'),'ok',first),
                       BatchItemResult('b',Path('同文件.wav'),'ok',second)))
    path=report.write_docx(batch,tmp_path/'共同记录.docx')
    contents=text(path)
    assert contents.count('共同唯一试件')==1
    assert '工况记录 C1' in contents and '处理记录 S1' in contents
    assert '[10, 200)' in contents and '第二事件' in contents
    assert contents.count('81.23456')==2


def test_batch_common_context_fields_once_and_per_item_differences_preserved(tmp_path):
    from dataclasses import replace
    first=replace(result(False),context=MeasurementContext(specimen='样件甲',supply='12 V',
                   motor_rotation='unknown',notes='方向未从文件名推断'))
    second=replace(first,result_id='second',context=replace(first.context,specimen='样件乙',
                   supply='14 V',actual_movement='向前',notes='操作者确认向前'))
    batch=BatchResult((BatchItemResult('a',first.path,'ok',first),BatchItemResult('b',second.path,'ok',second)))
    path=report.write_docx(batch,tmp_path/'共同基准差异.docx')
    contents=text(path)
    assert '工况记录 C1' in contents and '工况记录 C2' not in contents
    assert all(value in contents for value in ['样件甲','样件乙','12 V','14 V','向前','方向未从文件名推断','操作者确认向前'])
    cells=[cell.text for table in Document(path).tables for row in table.rows for cell in row.cells]
    assert cells.count('负载')==1 and cells.count('电机旋向')==1
    assert contents.count('81.23456')==2


def test_raw_head_metadata_replaced_by_digest_and_known_sensor_summary(tmp_path):
    from dataclasses import replace
    import hashlib,json
    from autoacoustics.model import to_plain
    metadata={'format':'HEAD_HDF','head':{'version':4,'release':6,'implementation_type':'FLOAT32',
       'fields':{'large_raw_header':'RAW_HEADER_MARKER '*4000},
       'comment_metadata':'_SI_Sensor = string(<SensorInfo Name="Mic001" Manufacturer="PCB" SerialNumber="73775" SensorType="35623" OutputUnit="mV" Sensitivity="50"><Quantity Unit="mV/Pa"/></SensorInfo>)'}}
    first=replace(result(False),provenance={'sample_rate':48000,'source_unit':'Pa','unit':'Pa','source_metadata':metadata})
    second=replace(first,result_id='second',provenance={**to_plain(first.provenance),
        'source_metadata':{**metadata,'different_raw_header':'UNIQUE_RAW_MARKER '*1000}})
    batch=BatchResult((BatchItemResult('a',first.path,'ok',first),BatchItemResult('b',second.path,'ok',second)))
    path=report.write_docx(batch,tmp_path/'HEAD 摘要.docx')
    contents=text(path)
    for sample in [first,second]:
        digest=hashlib.sha256(json.dumps(to_plain(sample.provenance['source_metadata']),ensure_ascii=False,
            sort_keys=True,separators=(',',':')).encode('utf-8')).hexdigest()
        assert digest in contents
    assert 'RAW_HEADER_MARKER' not in contents and 'UNIQUE_RAW_MARKER' not in contents
    assert 'HEAD_HDF' in contents and 'Pa' in contents and 'mV' in contents
    assert all(value in contents for value in ['Mic001','PCB','73775','35623','50'])
    assert contents.count('PCB')==1  # Same known sensor summary is shared.
    assert to_plain(first.provenance['source_metadata'])==metadata
    single=report.write_docx(first,tmp_path/'单文件 HEAD 摘要.docx')
    assert 'RAW_HEADER_MARKER' not in text(single)


def test_batch_item_headings_flow_without_forced_page_breaks(tmp_path):
    import xml.etree.ElementTree as ET
    batch=BatchResult((BatchItemResult('a',Path('a.wav'),'ok',result(False)),
                       BatchItemResult('b',Path('b.wav'),'ok',result(False,identity='second'))))
    path=report.write_docx(batch,tmp_path/'连续排版.docx')
    with zipfile.ZipFile(path) as archive:
        root=ET.fromstring(archive.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    headings=[p for p in root.findall('.//w:body/w:p',ns)
              if ''.join(t.text or '' for t in p.findall('.//w:t',ns)).startswith('条目 ')]
    assert len(headings)==2
    assert all(p.find('w:pPr/w:pageBreakBefore',ns) is None for p in headings)


def test_batch_method_and_quality_references_keep_actual_differences(tmp_path):
    from dataclasses import replace
    first=replace(result(False),summary=AnalysisSummary({'LZeq':MetricResult(81.,'dB',method='unique method alpha')},('来源质量甲',)))
    second=replace(first,result_id='second',summary=AnalysisSummary({'LZeq':MetricResult(82.,'dB',method='unique method beta')},('来源质量乙',)))
    batch=BatchResult((BatchItemResult('a',first.path,'ok',first),BatchItemResult('b',second.path,'ok',second)))
    path=report.write_docx(batch,tmp_path/'方法质量引用.docx')
    contents=text(path)
    assert '方法记录 M1' in contents and '方法记录 M2' in contents
    assert '质量记录 Q1' in contents and '质量记录 Q2' in contents
    assert contents.count('unique method alpha')==contents.count('unique method beta')==1
    assert contents.count('来源质量甲')==contents.count('来源质量乙')==1
    metric_rows=[row for table in Document(path).tables if len(table.columns)==5 for row in table.rows if 'LZeq' in row.cells[0].text]
    assert len(metric_rows)==2
    assert metric_rows[0].cells[4].text=='M1 / LZeq'
    assert metric_rows[1].cells[4].text=='M2 / LZeq'
